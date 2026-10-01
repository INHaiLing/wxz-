from io import BytesIO
from html.parser import HTMLParser
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.admin.models import LogEntry
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from tablib import Dataset

from content.importing import QuestionXLSX, SESSION_KEY
from content.models import Article, Category, Question
from content.resources import BASE_HEADERS, QuestionResource
from content.services import publish_questions


def make_dataset(*rows, answers=7, extra_headers=()):
    headers = [*BASE_HEADERS, *(f"答案{index + 1}" for index in range(answers)), *extra_headers]
    dataset = Dataset(headers=headers)
    for row in rows:
        dataset.append([row.get(header, "") for header in headers])
    return dataset


def literature_row(key="lit-001", **changes):
    row = {
        "题目ID": key, "来源": "literature", "分类ID": "qin", "篇目ID": "",
        "题型": "fact", "标签": "常识", "排序": 1,
        "题干": "《论语》记录{{0}}及其弟子的言行。", "答案1": "孔子",
    }
    row.update(changes)
    return row


class QuestionImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = get_user_model().objects.create_superuser("import-admin", "", "Local-test-password-2026")
        cls.category = Category.objects.create(id="qin", name="先秦")
        cls.article = Article.objects.create(id="quanxue", title="劝学")

    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.settings_override = override_settings(QUESTION_IMPORT_TMP_DIR=Path(self.temp.name))
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.client.force_login(self.owner)

    def preview(self, dataset, client=None):
        upload = SimpleUploadedFile(
            "questions.xlsx", QuestionXLSX().export_data(dataset),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return (client or self.client).post(reverse("admin:content_question_import"), {
            "import_file": upload, "format": "0", "resource": "0",
        })

    def confirm(self, preview, client=None):
        initial = dict(preview.context["confirm_form"].initial)
        return (client or self.client).post(reverse("admin:content_question_process_import"), initial)

    def create_question(self, key="lit-001"):
        return Question.objects.create(
            id=key, source="literature", category=self.category, type="fact",
            tag="常识", sort_order=1, stem="《论语》记录{{0}}及其弟子的言行。", answers=["孔子"],
        )

    def test_file_widget_has_text_companion_expected_by_unfold_upload_javascript(self):
        class WidgetTree(HTMLParser):
            def __init__(self):
                super().__init__()
                self.nodes = []
                self.stack = []

            def handle_starttag(self, tag, attrs):
                node = {"tag": tag, "attrs": dict(attrs), "parent": self.stack[-1] if self.stack else None}
                self.nodes.append(node)
                if tag not in {"input", "br", "img", "hr", "meta", "link"}:
                    self.stack.append(node)

            def handle_endtag(self, tag):
                if self.stack and self.stack[-1]["tag"] == tag:
                    self.stack.pop()

        response = self.client.get(reverse("admin:content_question_import"))
        self.assertEqual(response.status_code, 200)
        form = response.context["form"]
        tree = WidgetTree()
        tree.feed(form["import_file"].as_widget())
        upload = next(node for node in tree.nodes if node["attrs"].get("type") == "file")
        container = upload
        for _ in range(3):
            container = container["parent"]
        self.assertIsNotNone(container)
        companions = []
        for node in tree.nodes:
            if node["attrs"].get("type") != "text":
                continue
            ancestor = node["parent"]
            while ancestor is not None and ancestor is not container:
                ancestor = ancestor["parent"]
            if ancestor is container:
                companions.append(node)
        self.assertTrue(companions, "Unfold app.js requires a sibling text input three ancestors above the file input.")

    def test_template_has_blank_data_sheet_and_separate_help_examples(self):
        response = self.client.get(reverse("admin:content_question_import_template"))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["题库", "填写说明", "例题参考", "分类与篇目"])
        self.assertEqual(workbook["题库"].max_row, 1)
        self.assertEqual(workbook["题库"]["A1"].value, "题目ID")
        self.assertIn("{{0}}", workbook["例题参考"]["H2"].value)
        self.assertTrue(LogEntry.objects.filter(change_message__contains="下载题库").exists())
        self.assertEqual(len(QuestionXLSX().create_dataset(response.content)), 0)

    def test_real_xlsx_preview_confirmation_and_repeat_import(self):
        dataset = make_dataset(literature_row())
        preview = self.preview(dataset)
        self.assertEqual(preview.status_code, 200)
        self.assertContains(preview, "确认整批写入草稿")
        self.assertFalse(Question.objects.exists())
        self.assertEqual(preview.context["result"].totals["new"], 1)
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 302)
        question = Question.objects.get(pk="lit-001")
        self.assertEqual(question.answers, ["孔子"])
        self.assertFalse(question.is_published)
        self.assertIsNone(question.published_revision_id)
        self.assertTrue(LogEntry.objects.filter(object_id="lit-001").exists())
        self.assertEqual(list(Path(self.temp.name).iterdir()), [])
        repeat = self.preview(dataset)
        self.assertEqual(repeat.context["result"].totals["skip"], 1)
        self.assertEqual(self.confirm(repeat).status_code, 302)
        self.assertEqual(Question.objects.count(), 1)

    def test_mixed_invalid_batch_reports_row_and_rolls_back_every_row(self):
        dataset = make_dataset(literature_row("valid"), literature_row("invalid", 题干="没有占位符"))
        preview = self.preview(dataset)
        self.assertContains(preview, "校验错误 1")
        self.assertNotIn("confirm_form", preview.context)
        self.assertFalse(Question.objects.exists())
        result = QuestionResource(user=self.owner).import_data(dataset, dry_run=False)
        self.assertTrue(result.has_validation_errors())
        self.assertFalse(Question.objects.exists())

    def test_duplicate_and_missing_ids_are_row_errors_not_implicit_updates(self):
        dataset = make_dataset(literature_row("same"), literature_row("same"), literature_row(""))
        result = QuestionResource(user=self.owner).import_data(dataset, dry_run=False)
        self.assertEqual(len(result.invalid_rows), 3)
        self.assertFalse(Question.objects.exists())
        self.assertIn("文件内 ID 重复", str(result.invalid_rows[0].error))

    def test_import_cannot_set_publication_fields_or_write_partial_batch(self):
        dataset = make_dataset(literature_row(is_published=True), extra_headers=("is_published",))
        preview = self.preview(dataset)
        self.assertContains(preview, "不允许导入的列")
        self.assertNotIn("confirm_form", preview.context)
        self.assertFalse(Question.objects.exists())

    def test_import_updates_draft_without_changing_published_revision(self):
        question = self.create_question()
        publish_questions(Question.objects.filter(pk=question.pk), self.owner)
        question.refresh_from_db()
        previous_revision = question.published_revision_id
        preview = self.preview(make_dataset(literature_row(答案1="孔夫子")))
        self.assertEqual(preview.context["result"].totals["update"], 1)
        self.assertEqual(self.confirm(preview).status_code, 302)
        question.refresh_from_db()
        self.assertEqual(question.answers, ["孔夫子"])
        self.assertTrue(question.is_published)
        self.assertEqual(question.published_revision_id, previous_revision)
        self.assertEqual(question.published_revision.answers, ["孔子"])

    def test_preview_conflict_does_not_overwrite_new_edit_or_add_other_rows(self):
        question = self.create_question()
        preview = self.preview(make_dataset(literature_row(答案1="预览中的答案"), literature_row("other")))
        question.answers = ["后来编辑的答案"]
        question.save()
        response = self.confirm(preview)
        self.assertContains(response, "冲突 ID")
        question.refresh_from_db()
        self.assertEqual(question.answers, ["后来编辑的答案"])
        self.assertFalse(Question.objects.filter(pk="other").exists())

    def test_preview_of_new_id_cannot_overwrite_later_created_question(self):
        preview = self.preview(make_dataset(literature_row()))
        question = self.create_question()
        response = self.confirm(preview)
        self.assertContains(response, "重新上传预检")
        self.assertEqual(Question.objects.count(), 1)
        question.refresh_from_db()
        self.assertEqual(question.answers, ["孔子"])

    def test_permission_is_checked_on_upload_and_again_on_confirmation(self):
        user = get_user_model().objects.create_user("editor", password="Local-editor-test", is_staff=True)
        permissions = Permission.objects.filter(content_type__app_label="content", codename__in=("change_question", "import_question"))
        user.user_permissions.set(permissions)
        client = Client()
        client.force_login(user)
        no_add = self.preview(make_dataset(literature_row()), client=client)
        self.assertContains(no_add, "没有新增题目权限")
        self.create_question()
        preview = self.preview(make_dataset(literature_row(答案1="修改")), client=client)
        self.assertIn("confirm_form", preview.context)
        user.user_permissions.clear()
        self.assertEqual(self.confirm(preview, client=client).status_code, 403)
        denied = self.preview(make_dataset(literature_row()), client=client)
        self.assertEqual(denied.status_code, 403)

    def test_confirmation_is_bound_to_session_and_cannot_be_replayed(self):
        preview = self.preview(make_dataset(literature_row()))
        other_client = Client()
        other_client.force_login(self.owner)
        denied = self.confirm(preview, client=other_client)
        self.assertEqual(denied.status_code, 400)
        self.assertFalse(Question.objects.exists())
        self.assertEqual(self.confirm(preview).status_code, 302)
        self.assertEqual(self.confirm(preview).status_code, 400)
        self.assertEqual(Question.objects.count(), 1)
        self.assertFalse(self.client.session.get(SESSION_KEY))

    def test_expired_preview_and_changed_temporary_file_do_not_import(self):
        preview = self.preview(make_dataset(literature_row()))
        name = preview.context["confirm_form"].initial["import_file_name"]
        metadata_path = Path(self.temp.name) / f"{name}.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        metadata["created"] -= 3600
        metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
        self.assertEqual(self.confirm(preview).status_code, 400)
        self.assertFalse(Question.objects.exists())
        preview = self.preview(make_dataset(literature_row()))
        name = preview.context["confirm_form"].initial["import_file_name"]
        (Path(self.temp.name) / f"{name}.bin").write_bytes(b"changed after preview")
        self.assertEqual(self.confirm(preview).status_code, 400)
        self.assertFalse(Question.objects.exists())

    def test_export_roundtrip_preserves_id_dynamic_answers_and_literal_formula_text(self):
        answers = [f"答案{index}" for index in range(10)]
        stem = "=这是一段文字，" + "、".join("{{%d}}" % index for index in range(10))
        question = Question.objects.create(
            id="many-answers", source="classical", article=self.article, type="translation",
            stem=stem, answers=answers,
        )
        response = self.client.post(reverse("admin:content_question_export"), {"format": "0", "resource": "0"})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        sheet = workbook["题库"]
        self.assertEqual(sheet["A2"].value, question.pk)
        self.assertEqual(sheet["H2"].value, stem)
        self.assertEqual(sheet["H2"].data_type, "s")
        self.assertEqual(sheet.max_column, len(BASE_HEADERS) + 10)
        dataset = QuestionXLSX().create_dataset(response.content)
        preview = self.preview(dataset)
        self.assertEqual(preview.context["result"].totals["skip"], 1)
        self.assertEqual(self.confirm(preview).status_code, 302)
        question.refresh_from_db()
        self.assertEqual(question.answers, answers)

    def test_xlsx_formula_size_and_archive_limits_are_rejected(self):
        raw = QuestionXLSX().export_data(make_dataset(literature_row()))
        workbook = load_workbook(BytesIO(raw))
        workbook.active["H2"] = '=HYPERLINK("https://example.com")'
        output = BytesIO()
        workbook.save(output)
        with self.assertRaisesRegex(ValueError, "包含公式"):
            QuestionXLSX().create_dataset(output.getvalue())
        with patch("content.importing.MAX_UNCOMPRESSED_BYTES", 10):
            with self.assertRaisesRegex(ValueError, "解压后过大"):
                QuestionXLSX().create_dataset(raw)
        with patch("content.importing.MAX_IMPORT_ROWS", 1):
            two = QuestionXLSX().export_data(make_dataset(literature_row("one"), literature_row("two")))
            with self.assertRaisesRegex(ValueError, "最多 1 行"):
                QuestionXLSX().create_dataset(two)
        with patch("content.importing.MAX_UPLOAD_BYTES", 10):
            with self.assertRaisesRegex(ValueError, "超过 5 MiB"):
                QuestionXLSX().create_dataset(raw)

    def test_invalid_ownership_and_answer_gaps_are_rejected_by_model_validation(self):
        dataset = make_dataset(
            literature_row("bad-owner", 篇目ID="quanxue"),
            literature_row("gap", 题干="{{0}}、{{1}}、{{2}}", 答案1="甲", 答案2="", 答案3="丙"),
        )
        result = QuestionResource(user=self.owner).import_data(dataset, dry_run=False)
        self.assertEqual(len(result.invalid_rows), 2)
        self.assertFalse(Question.objects.exists())

    def test_large_answer_column_number_is_bounded_before_allocating_range(self):
        row = literature_row()
        row["答案999999999"] = "不合法列"
        preview = self.preview(make_dataset(row, extra_headers=("答案999999999",)))
        self.assertContains(preview, "答案列编号不能超过 32")
        self.assertNotIn("confirm_form", preview.context)
        self.assertFalse(Question.objects.exists())

    def test_blank_middle_row_keeps_excel_error_line_and_bad_reference_is_readable(self):
        dataset = make_dataset(literature_row("valid"), {}, literature_row("bad-fk", 分类ID="does-not-exist"))
        preview = self.preview(dataset)
        self.assertContains(preview, "Excel 第 3 行")
        self.assertContains(preview, "Excel 第 4 行")
        self.assertContains(preview, "找不到文学分类 ID")
        self.assertFalse(Question.objects.exists())
