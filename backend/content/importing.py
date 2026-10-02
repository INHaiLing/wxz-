"""XLSX transport and secure two-step confirmation around django-import-export."""

from io import BytesIO
from pathlib import Path
import hashlib
import json
import re
import time
import uuid
from zipfile import BadZipFile, ZipFile

from django import forms
from django.conf import settings
from django.contrib import messages
from django.contrib.admin.models import CHANGE, LogEntry
from django.contrib.contenttypes.models import ContentType
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.template.response import TemplateResponse
from django.urls import path, reverse
from django.utils.decorators import method_decorator
from django.views.decorators.http import require_POST
from import_export.formats.base_formats import XLSX
from import_export.forms import ConfirmImportForm
from import_export.tmp_storages import BaseStorage
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tablib import Dataset
from unfold.contrib.import_export.forms import ExportForm, ImportForm

from .models import Article, Category
from .resources import (
    BASE_HEADERS, DEFAULT_ANSWER_COLUMNS, MAX_ANSWER_COLUMNS, MAX_IMPORT_ROWS,
    import_permission,
)


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
PREVIEW_SECONDS = 30 * 60
SESSION_KEY = "question_import_previews"
SIGNING_SALT = "content.question-import.v1"
SAFE_NAME = re.compile(r"[0-9a-f]{32}\Z")


def _storage_directory():
    return Path(getattr(settings, "QUESTION_IMPORT_TMP_DIR", Path(settings.BASE_DIR) / ".local" / "imports")).resolve()


class QuestionTemporaryStorage(BaseStorage):
    """Random, private files; user-provided paths can never select another file."""

    def get_full_path(self):
        if not self.name or not SAFE_NAME.fullmatch(self.name):
            raise ValidationError("导入凭证无效，请重新上传文件。")
        root = _storage_directory()
        target = (root / f"{self.name}.bin").resolve()
        if target.parent != root:
            raise ValidationError("导入临时文件路径无效。")
        return target

    def save(self, data):
        self.name = uuid.uuid4().hex
        target = self.get_full_path()
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("xb") as handle:
            handle.write(data)

    def read(self):
        target = self.get_full_path()
        if target.stat().st_size > MAX_UPLOAD_BYTES:
            raise ValidationError("上传文件超过 5 MiB，请分批导入。")
        return target.read_bytes()

    def remove(self):
        target = self.get_full_path()
        target.unlink(missing_ok=True)
        target.with_suffix(".json").unlink(missing_ok=True)


def _write_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    for row in [headers, *rows]:
        sheet.append(list(row))
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                # XLSX stores a real string, not a formula or an apostrophe-mutated value.
                cell.data_type = "s"
                cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="175C52")
        cell.font = Font(color="FFFFFF", bold=True)
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 52 if header in ("题干", "说明", "示例") else 20
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


class QuestionXLSX(XLSX):
    def create_dataset(self, in_stream, **kwargs):
        raw = bytes(in_stream)
        if len(raw) > MAX_UPLOAD_BYTES:
            raise ValueError("上传文件超过 5 MiB，请分批导入。")
        try:
            with ZipFile(BytesIO(raw)) as archive:
                files = archive.infolist()
                if len(files) > 2000 or sum(entry.file_size for entry in files) > MAX_UNCOMPRESSED_BYTES:
                    raise ValueError("XLSX 解压后过大，请删除多余工作表、格式或分批导入。")
                if any(entry.flag_bits & 1 for entry in files):
                    raise ValueError("不支持加密或带密码的 XLSX 文件。")
                if "[Content_Types].xml" not in archive.namelist():
                    raise ValueError("不是有效的 XLSX 工作簿。")
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        except (BadZipFile, OSError, KeyError) as exc:
            raise ValueError("无法读取 XLSX 文件，请使用下载的模板并另存为 .xlsx。") from exc
        try:
            sheet = workbook["题库"] if "题库" in workbook.sheetnames else workbook.worksheets[0]
            max_columns = len(BASE_HEADERS) + MAX_ANSWER_COLUMNS
            if sheet.max_column and sheet.max_column > max_columns:
                raise ValueError(f"题库工作表最多 {max_columns} 列，请删除额外列或隐藏发布字段。")
            if sheet.max_row and sheet.max_row > MAX_IMPORT_ROWS + 1:
                raise ValueError(f"题库工作表最多 {MAX_IMPORT_ROWS} 行数据，请分批导入并删除多余空行格式。")
            rows = sheet.iter_rows()
            first = next(rows, ())
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in first]
            while headers and not headers[-1]:
                headers.pop()
            dataset = Dataset(headers=headers)
            pending_empty_rows = []
            for excel_line, cells in enumerate(rows, 2):
                if excel_line > MAX_IMPORT_ROWS + 1:
                    raise ValueError(f"题库工作表最多 {MAX_IMPORT_ROWS} 行数据。")
                values = [cell.value for cell in cells[:len(headers)]]
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Excel 第 {excel_line} 行包含公式，请粘贴为纯文本或数值后导入。")
                if any(cell.value is not None for cell in cells[len(headers):]):
                    raise ValueError(f"Excel 第 {excel_line} 行存在没有表头的数据列，请补齐合法列名。")
                if any(value is not None for value in values):
                    for empty_row in pending_empty_rows:
                        dataset.append(empty_row)
                    pending_empty_rows.clear()
                    dataset.append(values)
                else:
                    pending_empty_rows.append(values)
            return dataset
        finally:
            workbook.close()

    def export_data(self, dataset, **kwargs):
        workbook = Workbook()
        workbook.remove(workbook.active)
        _write_sheet(workbook, "题库", dataset.headers, list(dataset))
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()


class QuestionImportForm(ImportForm):
    def clean_import_file(self):
        upload = self.cleaned_data["import_file"]
        if not upload.name.lower().endswith(".xlsx"):
            raise forms.ValidationError("仅支持 .xlsx 文件，请先下载题库模板。")
        if not 0 < upload.size <= MAX_UPLOAD_BYTES:
            raise forms.ValidationError("文件不能为空，且不能超过 5 MiB。")
        return upload


class QuestionConfirmImportForm(ConfirmImportForm):
    confirmation_token = forms.CharField(widget=forms.HiddenInput())


class QuestionImportMixin:
    import_template_name = "admin/content/question/import.html"
    import_form_class = QuestionImportForm
    confirm_form_class = QuestionConfirmImportForm
    # Always export the complete round-trip schema, including stable IDs and all answers.
    export_form_class = ExportForm
    skip_admin_log = False

    def get_import_formats(self):
        return [QuestionXLSX]

    def get_export_formats(self):
        return [QuestionXLSX]

    def get_tmp_storage_class(self):
        return QuestionTemporaryStorage

    def is_skip_import_confirm_enabled(self):
        return False

    def has_import_permission(self, request):
        return import_permission(request.user)

    def has_export_permission(self, request):
        return bool(request.user.is_active and request.user.is_staff and self.has_view_permission(request))

    def get_urls(self):
        info = self.get_model_info()
        return [path(
            "import-template/", self.admin_site.admin_view(self.download_import_template),
            name="%s_%s_import_template" % info,
        )] + super().get_urls()

    def get_import_resource_kwargs(self, request, **kwargs):
        values = super().get_import_resource_kwargs(request, **kwargs)
        values.update(user=request.user, request=request)
        if hasattr(request, "_question_confirm_versions"):
            values["expected_versions"] = request._question_confirm_versions
        return values

    def get_confirm_form_initial(self, request, import_form):
        initial = super().get_confirm_form_initial(request, import_form)
        if import_form is None:
            return initial
        storage = QuestionTemporaryStorage(name=initial["import_file_name"])
        digest = hashlib.sha256(storage.read()).hexdigest()
        metadata = {
            "owner": str(request.user.pk), "created": time.time(), "digest": digest,
            "versions": request._question_import_versions,
        }
        storage.get_full_path().with_suffix(".json").write_text(
            json.dumps(metadata, ensure_ascii=False), encoding="utf-8",
        )
        token = signing.dumps({"name": storage.name, "owner": str(request.user.pk), "digest": digest}, salt=SIGNING_SALT)
        pending = request.session.get(SESSION_KEY, {})
        while len(pending) >= 5:
            oldest = next(iter(pending))
            pending.pop(oldest)
            QuestionTemporaryStorage(name=oldest).remove()
        pending[storage.name] = hashlib.sha256(token.encode()).hexdigest()
        request.session[SESSION_KEY] = pending
        initial["confirmation_token"] = token
        return initial

    def _discard_preview(self, request, storage):
        pending = request.session.get(SESSION_KEY, {})
        pending.pop(storage.name, None)
        request.session[SESSION_KEY] = pending
        storage.remove()

    def _cleanup_expired(self):
        root = _storage_directory()
        if root.exists():
            for target in root.glob("*.bin"):
                if SAFE_NAME.fullmatch(target.stem) and target.resolve().parent == root:
                    try:
                        if time.time() - target.stat().st_mtime > PREVIEW_SECONDS:
                            QuestionTemporaryStorage(name=target.stem).remove()
                    except FileNotFoundError:
                        pass  # Another completed request may already have removed it.

    def _context(self, request, **extra):
        context = self.admin_site.each_context(request)
        context.update(
            title="题库草稿导入", opts=self.model._meta, media=self.media,
            import_template_url=reverse("admin:%s_%s_import_template" % self.get_model_info()),
            import_url=reverse("admin:%s_%s_import" % self.get_model_info()),
            import_error_display=("message",), **extra,
        )
        return context

    def import_action(self, request, **kwargs):
        if not self.has_import_permission(request):
            raise PermissionDenied
        self._cleanup_expired()
        response = super().import_action(request, **kwargs)
        if hasattr(response, "context_data"):
            response.context_data.update({
                "import_template_url": reverse("admin:%s_%s_import_template" % self.get_model_info()),
                "import_url": reverse("admin:%s_%s_import" % self.get_model_info()),
            })
            form = response.context_data.get("form")
            upload = form.cleaned_data.get("import_file") if form and hasattr(form, "cleaned_data") else None
            if upload is not None and hasattr(upload, "tmp_storage_name") and not response.context_data.get("confirm_form"):
                self._discard_preview(request, QuestionTemporaryStorage(name=upload.tmp_storage_name))
        return response

    def add_data_read_fail_error_to_form(self, form, error):
        message = str(error) if isinstance(error, ValueError) else "无法解析工作簿，请检查文件是否损坏，并重新使用 XLSX 模板。"
        form.add_error("import_file", message)

    @method_decorator(require_POST)
    def process_import(self, request, **kwargs):
        if not self.has_import_permission(request):
            raise PermissionDenied
        form = self.create_confirm_form(request)
        storage = None
        try:
            if not form.is_valid():
                raise ValidationError("导入确认信息不完整，请重新上传预检。")
            token = form.cleaned_data["confirmation_token"]
            signed = signing.loads(token, salt=SIGNING_SALT, max_age=PREVIEW_SECONDS)
            pending = request.session.get(SESSION_KEY, {})
            name = form.cleaned_data["import_file_name"]
            if (
                signed.get("name") != name or signed.get("owner") != str(request.user.pk)
                or pending.get(name) != hashlib.sha256(token.encode()).hexdigest()
                or form.cleaned_data["format"] != "0"
                or form.cleaned_data.get("resource") not in ("", "0")
            ):
                raise ValidationError("导入凭证不属于当前会话，或已经使用，请重新上传。")
            storage = QuestionTemporaryStorage(name=name)
            raw = storage.read()
            metadata = json.loads(storage.get_full_path().with_suffix(".json").read_text(encoding="utf-8"))
            if (
                metadata["owner"] != str(request.user.pk)
                or time.time() - metadata["created"] > PREVIEW_SECONDS
                or hashlib.sha256(raw).hexdigest() != signed["digest"]
                or metadata["digest"] != signed["digest"]
            ):
                raise ValidationError("预览已失效或文件发生变化，请重新上传。")
            dataset = QuestionXLSX().create_dataset(raw)
            request._question_confirm_versions = metadata["versions"]
            with transaction.atomic():
                result = super().process_dataset(dataset, form, request, **kwargs)
                failed = result.has_errors() or result.has_validation_errors()
                if failed:
                    transaction.set_rollback(True)
                else:
                    # Import-export creates Django LogEntry records within this same transaction.
                    response = self.process_result(result, request)
            if failed:
                # Admin navigation may query singleton configuration modules.
                # Build it only after the failed batch's savepoint has rolled back.
                response = TemplateResponse(request, self.import_template_name, self._context(
                    request, result=result, confirmation_error="本批未写入任何题目。请按错误提示修正后重新上传预检。",
                ))
            self._discard_preview(request, storage)
            return response
        except (ValidationError, signing.BadSignature, FileNotFoundError, ValueError, KeyError) as exc:
            if storage is not None:
                self._discard_preview(request, storage)
            message = "；".join(exc.messages) if isinstance(exc, ValidationError) else "预览已过期、文件不可用或确认凭证无效，请重新上传预检。"
            return TemplateResponse(request, self.import_template_name, self._context(request, confirmation_error=message), status=400)

    def add_success_message(self, result, request):
        messages.success(request, "草稿导入完成：新增 %(new)s、修改 %(update)s、无变化 %(skip)s。题目尚未发布。" % result.totals)

    def get_export_data(self, file_format, request, queryset, **kwargs):
        if not self.has_export_permission(request):
            raise PermissionDenied
        output = super().get_export_data(file_format, request, queryset, **kwargs)
        self._log_file_action(request, "导出题库草稿 XLSX（保留稳定 ID）")
        return output

    def _log_file_action(self, request, message):
        LogEntry.objects.create(
            user_id=request.user.pk, content_type=ContentType.objects.get_for_model(self.model),
            object_id="", object_repr="题库导入导出", action_flag=CHANGE, change_message=message,
        )

    def download_import_template(self, request):
        if not self.has_import_permission(request):
            raise PermissionDenied
        workbook = Workbook()
        workbook.remove(workbook.active)
        headers = [*BASE_HEADERS, *(f"答案{index + 1}" for index in range(DEFAULT_ANSWER_COLUMNS))]
        _write_sheet(workbook, "题库", headers, [])
        instructions = [
            ("操作", "只填写首个“题库”工作表。其他页仅供说明，不参与导入。先预检，再确认写入草稿，最后在后台发布。"),
            ("题目ID", "必须填写稳定唯一 ID，并将单元格格式设为文本。已有 ID 表示修改，新 ID 表示新增；切勿删除 ID 后重新导入。"),
            ("来源与题型", "literature=文学常识（仅 fact）；classical=文言文，题型可为 fact、word 或 translation。"),
            ("归属", "文学题填分类ID、篇目ID留空；文言题填篇目ID、分类ID留空。使用“分类与篇目”表中的现有 ID。"),
            ("题干与答案", "题干用 {{0}}、{{1}}…标记答案；对应答案1、答案2…列。至少1个答案，中间不能空列；可重复引用同一编号。"),
            ("答案列", "模板含7列；需要更多答案时连续增加答案8…答案32。尾部不用的答案格留空。"),
            ("限制", "仅 XLSX，文件不超过5 MiB、题库不超过5000行、解压后不超过50 MiB；不支持公式、密码、Word、PDF或图片识别。"),
            ("安全与冲突", "不得加入发布状态、版本、权限等列。预览30分钟有效；其间题目被修改/发布/下架，需重新预检。"),
            ("重复文件", "相同ID且内容未变的行显示“无变化”，不会重复新增。任一行错误时整批不写入。"),
        ]
        _write_sheet(workbook, "填写说明", ("项目", "说明"), instructions)
        category = Category.objects.order_by("sort_order", "pk").first()
        article = Article.objects.order_by("sort_order", "pk").first()
        examples = [
            ["example-lit-001", "literature", category.pk if category else "请先新建分类", "", "fact", "常识", 1, "《论语》记录了{{0}}及其弟子的言行。", "孔子"],
            ["example-word-001", "classical", "", article.pk if article else "请先新建篇目", "word", "重点字词", 2, "“青，取之于蓝，而青于蓝”中，“于”意为{{0}}。", "比"],
        ]
        _write_sheet(workbook, "例题参考", headers, [row + [""] * (len(headers) - len(row)) for row in examples])
        references = [("分类", obj.pk, obj.name, "启用" if obj.is_active else "停用（可保存草稿，不可发布）") for obj in Category.objects.order_by("sort_order", "pk")]
        references += [("篇目", obj.pk, obj.title, "启用" if obj.is_active else "停用（可保存草稿，不可发布）") for obj in Article.objects.order_by("sort_order", "pk")]
        _write_sheet(workbook, "分类与篇目", ("类型", "ID", "名称", "状态"), references)
        output = BytesIO()
        workbook.save(output)
        self._log_file_action(request, "下载题库 XLSX 模板")
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="question-template.xlsx"'
        return response
